#
# Compare the DMA Repo Extract Assets to the DMA Assessment Spreadsheet
# Duplicate names are going to be missing in the repo as the name must be unique
# Anything else may have a name normalization issue, or it may have actually been missed
#

import json
import openpyxl

DMA_REPO_EXTRACT_FILE = '$DMA_Repo_Extract.json'
DMA_ASSET_SPREADSHEET = '$LOE_Assessment_dma_splunk_53_com_2026-07-24_MODIFIED.xlsx'

def main():
    dma_repo_extract_document_name_type_pairs = set()
    with open(DMA_REPO_EXTRACT_FILE, 'r', encoding='utf-8') as infile:
        input_string = infile.read()
        input_json = parse_json_safely(input_string)
        name = input_json.get('name')
        documents = input_json.get('documents')
        for document in documents:
            document_name = document.get('document_name')
            document_type = document.get('document_type')
            # print(document_name, document_type)
            dma_repo_extract_document_name_type_pairs.add((get_normalized_asset_name(document_name), document_type))

    print(dma_repo_extract_document_name_type_pairs)

    wb = openpyxl.load_workbook(DMA_ASSET_SPREADSHEET, data_only=True)
    sheet = wb['Dashboards']

    for row in sheet.iter_rows(min_row=2, values_only=True):
        # CAUTION: Dashboard Label is used for the name
        asset_name = str(row[2]) if row[2] is not None else ""
        asset_status = str(row[18]) if row[18] is not None else ""
        # print(f"Asset Name: {asset_name} | Asset Status: {asset_status}")
        if asset_status != 'Backlog':
            normalized_asset_name = get_normalized_asset_name(asset_name)
            if ((normalized_asset_name, 'dashboard')) not in dma_repo_extract_document_name_type_pairs:
                print(f'Dashboard "{normalized_asset_name} ({asset_name})" failed coverage check!')

    wb = openpyxl.load_workbook(DMA_ASSET_SPREADSHEET, data_only=True)
    sheet = wb['Searches & Alerts']

    for row in sheet.iter_rows(min_row=2, values_only=True):
        asset_name = str(row[1]) if row[1] is not None else ""
        asset_type = str(row[2]) if row[2] is not None else ""
        asset_status = str(row[14]) if row[14] is not None else ""
        # print(f"Asset Name: {asset_name} | Asset Status: {asset_status}")
        if asset_status != 'Backlog':
            normalized_asset_name = get_normalized_asset_name(asset_name)
            augmented_asset_name = get_augmented_asset_name(normalized_asset_name, asset_type)
            if ((augmented_asset_name, 'notebook')) not in dma_repo_extract_document_name_type_pairs:
                print(f'Notebook "{augmented_asset_name} ({asset_name}) [{asset_type}]" failed coverage check!')



def get_augmented_asset_name(asset_name, asset_type):
    prefix_dict = {
        'alert': '[ALERT] - ',
        'report': '[REPORT] - ',
        'Saved Search': '[SEARCH] - ',
    }

    prefix = prefix_dict.get(asset_type, '')

    return prefix + asset_name


def get_normalized_asset_name(asset_name):
    normalized_asset_name = asset_name.replace('[Splunk] ', '')
    normalized_asset_name = normalized_asset_name.replace(':', '-')
    normalized_asset_name = normalized_asset_name.replace('/', '-')
    normalized_asset_name = normalized_asset_name.replace('&amp;', '&')
    normalized_asset_name = normalized_asset_name.replace('&gt;', '-')
    normalized_asset_name = normalized_asset_name.replace('>', '-')
    normalized_asset_name = normalized_asset_name.replace('<', '-')
    normalized_asset_name = normalized_asset_name.replace('|', '-')
    normalized_asset_name = normalized_asset_name.replace('"', '-')
    normalized_asset_name = normalized_asset_name.replace('%', '-')
    normalized_asset_name = normalized_asset_name.replace('?', '-')
    normalized_asset_name = normalized_asset_name.replace('*', '-')
    normalized_asset_name = normalized_asset_name.strip()


    return normalized_asset_name

def parse_json_safely(json_string):
    try:
        data = json.loads(json_string)
        # print("✅ Success! Parsed data:", data)
        return data

    except json.JSONDecodeError as e:
        # Handles malformed JSON strings (e.g., missing quotes, trailing commas)
        print(f"❌ JSON Decode Error: Invalid JSON syntax.")
        print(f"Details: {e.msg} at line {e.lineno}, column {e.colno}")
        return None

    except TypeError as e:
        # Handles incorrect data types passed to loads() (e.g., passing None or an integer)
        print(f"❌ Type Error: Expected a string-like object, but got something else. Details: {e}")
        return None


if __name__ == '__main__':
    main()
